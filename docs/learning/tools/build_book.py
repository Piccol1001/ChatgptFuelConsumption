#!/usr/bin/env python3
"""Build the learning editions from Markdown and an immutable Git source snapshot.

Documentation tools only: Python 3, pandoc, reportlab, python-docx, pypdf.
Run from any directory: python docs/learning/tools/build_book.py
No Android sources, app versions, signing material or user data are modified.
"""
from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import re
import subprocess
import textwrap
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, KeepTogether, PageBreak, PageTemplate,
    Paragraph, Preformatted, Spacer, Table, TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parents[3]
DOCS = ROOT / "docs/learning"
OUT = ROOT / "learning-materials"
SNAPSHOT = "1f13482bfc550da77e5c950e48075f62de12620f"
SOURCE_URL = f"https://github.com/Piccol1001/ChatgptFuelConsumption/blob/{SNAPSHOT}/"
TITLE = "Fuel & Expense Tracker — Learning Guide"
BASE = "Fuel_Expense_Tracker_Learning_Guide"
SOURCE_PREFIX = "app/src/main/java/my/fuelledger/"

DESCRIPTIONS = {
    "LedgerApplication.kt": "Application-owned vault, database and repository composition.",
    "MainActivity.kt": "Android entry, lifecycle, authentication gate and Compose root.",
    "LedgerViewModel.kt": "Ledger state, receipt operations, preferences, backup and UI events.",
    "FinanceViewModel.kt": "Income, recurring rules, budgets and monthly planning state.",
    "JourneyViewModel.kt": "Saved journey state and repository operations.",
    "data/Backup.kt": "Backup model validation, relationships and safe CSV encoding.",
    "data/BackupCrypto.kt": "Password-derived authenticated streaming backup encryption.",
    "data/Database.kt": "Room entities, DAOs, constraints, schema and migrations 1–5.",
    "data/ExpenseEntry.kt": "Payment display grouping and manual category availability.",
    "data/ExpenseTools.kt": "Favorites, Undo models and expense-tool validation.",
    "data/FinanceData.kt": "Income, budgets, recurring models, validation and anchored dates.",
    "data/FinanceRepository.kt": "Recurring reconciliation, commitment payment and planning persistence.",
    "data/JourneyData.kt": "Versioned journey payloads, estimates, actual links and storage.",
    "data/LedgerRepository.kt": "Atomic ledger writes, duplicate safeguards, Undo and restore.",
    "data/MonthlyPlanner.kt": "Planned occurrences, allowance and exact allocation calculations.",
    "data/ReceiptAutomation.kt": "Merchant memory, fingerprints and conservative auto-save gates.",
    "data/ReceiptBackupArchive.kt": "Bounded receipt archive staging, validation and restoration.",
    "data/Refinement.kt": "Stable category order, fuel month series and flexible allocation movement.",
    "data/SettingsStore.kt": "Typed encrypted DataStore and legacy settings migration.",
    "domain/Calculations.kt": "Exact money, measurements, trip estimates, fuel intervals and spending.",
    "domain/CentsKeypad.kt": "Bounded integer digit entry and backspace.",
    "domain/FinancialReport.kt": "Actual income/expense reports and local financial insights.",
    "domain/VehicleAnalysis.kt": "Coverage-aware vehicle cost, weighted price and consumption changes.",
    "finance/PlannerReminders.kt": "Local planner notification generation.",
    "finance/RecurringWorker.kt": "Unique periodic reconciliation via WorkManager.",
    "receipts/ReceiptParser.kt": "Pure text candidate extraction for Malaysian receipt formats.",
    "receipts/ReceiptScanner.kt": "Bundled ML Kit recognition, image quality and resource lifetime.",
    "receipts/ReceiptStore.kt": "Private encrypted photos, normalization, crop, rotate and limits.",
    "security/AppLock.kt": "Device authentication and session/lifecycle lock policy.",
    "security/EncryptedDatabaseFiles.kt": "Protected database passphrases and preserving conversion.",
    "security/LocalVault.kt": "Keystore AES-GCM encryption with purpose binding.",
    "security/PreservingOpenHelperFactory.kt": "SQLCipher integration that preserves files on open failure.",
    "ui/AppLockSettings.kt": "Authentication gate and lock preference interface.",
    "ui/Components.kt": "Reusable cards, headings, fields, icons and shared UI helpers.",
    "ui/DailySpendingChart.kt": "Scrollable daily amounts, date selection and Today positioning.",
    "ui/DesignTokens.kt": "Shared spacing, shape, target and dashboard color constants.",
    "ui/Editors.kt": "Existing expense/refuel and supporting editing flows.",
    "ui/ExpenseEntryControls.kt": "Category icons, payment buttons and cents keypad UI.",
    "ui/ExpenseToolsUi.kt": "Favorites, refund/Undo and transaction-tool controls.",
    "ui/FinanceScreen.kt": "Income, budgets and recurring transaction controls.",
    "ui/FinancialAnalysisUi.kt": "Financial planning/actual summaries and report presentation.",
    "ui/FintechTheme.kt": "Material palettes, typography and theme selection.",
    "ui/JourneyScreen.kt": "Journey form, estimates and saved journey presentation.",
    "ui/JourneyToolsUi.kt": "Route tools, actual-link selection and sharing.",
    "ui/LedgerApp.kt": "Primary Scaffold, navigation and screen/form orchestration.",
    "ui/LocalDataSettings.kt": "Local retention/cleanup controls.",
    "ui/MainPager.kt": "Main destinations, HorizontalPager and gesture boundaries.",
    "ui/MonthlyPlannerScreen.kt": "Monthly planning and allocation editing interface.",
    "ui/OverviewDashboard.kt": "Home overview, budget, quick actions and upcoming items.",
    "ui/PersonalizationSettings.kt": "Preference UI and shared UI preference context.",
    "ui/PressComponents.kt": "Material component adapters for consistent interaction feedback.",
    "ui/PressFeedback.kt": "Draw-only press scaling, rounded focus and motion settings.",
    "ui/PrivacyText.kt": "Privacy-aware value presentation.",
    "ui/ReceiptAttachment.kt": "Receipt attachment, preview and image transformation UI.",
    "ui/ReceiptCameraButton.kt": "Camera intent/FileProvider capture integration.",
    "ui/RefinedControls.kt": "Shared calendar and refined supporting controls.",
    "ui/ReportActions.kt": "Explicit local CSV/PDF document exports.",
    "ui/ScreenPreviews.kt": "Compose preview entry points and fixture states.",
    "ui/Screens.kt": "Main fuel, expense and analysis screen implementations.",
    "ui/SettingsScreen.kt": "Settings categories, scanner mode, About and local workflows.",
    "ui/SmartReceiptScreen.kt": "Extracted receipt review and acceptance UI.",
    "ui/VehicleAnalysisUi.kt": "Coverage and vehicle comparison presentation.",
    "widgets/LedgerWidgets.kt": "Android widgets, privacy-aware values and quick actions.",
}


def git_bytes(path: str) -> bytes:
    return subprocess.check_output(["git", "show", f"{SNAPSHOT}:{path}"], cwd=ROOT)


def pinned_link(path: str, line: int | None = None) -> str:
    return SOURCE_URL + path + (f"#L{line}" if line else "")


def source_snippet(title: str, path: str, start: str, end: str | None, explanation: str) -> str:
    lines = git_bytes(path).decode().splitlines()
    first = next(i for i, line in enumerate(lines) if line.startswith(start))
    last = next((i for i in range(first + 1, len(lines)) if end and lines[i].startswith(end)), len(lines))
    excerpt = "\n".join(lines[first:last]).rstrip()
    return (f"## {title}\n\n{explanation}\n\n"
            f"Exact source: [{path}, line {first + 1}]({pinned_link(path, first + 1)}). "
            "Imports or surrounding declarations may be outside this excerpt; use the companion for the complete file.\n\n"
            f"```kotlin\n{excerpt}\n```\n\n")


def appendices(paths: list[str]) -> str:
    p = SOURCE_PREFIX
    parts = ["# Appendix A. Read real code, one decision at a time\n\n"
             "These excerpts are extracted directly from the pinned Git checkpoint. "
             "They are not newly invented replacement implementations. Read the explanation, "
             "predict a result, then compare with the function and its tests.\n\n"]
    snippets = [
        ("A1. Exact money", "domain/Calculations.kt", "object Money", "fun measurement",
         "The regex limits manual input to decimal money; moving the decimal point converts RM to sen. "
         "longValueExact rejects non-integral or overflowing conversion. The range check rejects zero and excess. "
         "Estimates deliberately use rounding; formatting never changes the stored integer."),
        ("A2. The entire keypad algorithm", "domain/CentsKeypad.kt", "fun appendSenDigit", None,
         "require validates the function's input contract. The rearranged inequality checks the maximum before multiplication. "
         "A null result means reject the proposed digit. Integer division drops the final digit."),
        ("A3. A test that protects the money boundary", "../test/java/my/fuelledger/domain/CentsKeypadTest.kt",
         "class CentsKeypadTest", None,
         "The tests cover the user's exact keystroke sequence, editing compatibility and inclusive maximum. "
         "Expected formatted values are written explicitly, rather than recomputed with the same formula."),
        ("A4. Fuel intervals", "domain/Calculations.kt", "fun fillEstimates", "private fun validFuelPoint",
         "sortBy establishes chronological order. zipWithNext creates adjacent pairs. mapNotNull discards invalid intervals. "
         "The later entry supplies purchased litres and cost; the odometer difference supplies distance. "
         "This implements the estimate and its assumptions discussed in Chapter 11."),
        ("A5. A refuel and one linked expense", "data/LedgerRepository.kt", "    suspend fun saveRefuel", "    suspend fun deleteRefuel",
         "The repository lock coordinates its writes. Validation precedes changes. Inside the Room transaction, "
         "the edited point is checked against the entire sequence. Existing linked identity/payment are preserved, "
         "and both rows commit together. The edit's before/after records support Undo."),
        ("A6. Historical payment preservation", "data/ExpenseEntry.kt", "val expensePaymentOptions", "/** Hide",
         "The visible Card group accepts several old stored strings. If the selected group is unchanged, "
         "the original string is returned. Grouping a display is not the same as migrating historical data."),
        ("A7. Fixed commitments are excluded once", "data/MonthlyPlanner.kt", "fun variableExpensesForMonth", "data class MonthlyMoney",
         "The reserved set identifies rule/date occurrences, while fixedIds identifies their actual payments. "
         "The filter selects the actual month and excludes those reserved payments and their refunds. "
         "An unreserved payment has different allowance treatment."),
        ("A8. Largest-remainder allocation", "data/MonthlyPlanner.kt", "fun allocationAmounts", "@Dao",
         "BigInteger avoids overflowing the intermediate multiplication. Quotients give floor amounts. "
         "Sorted remainders determine who gets each leftover sen. Savings has no spending-category expense deduction."),
        ("A9. Flexible percentage movement", "data/Refinement.kt", "fun moveAllocation", "fun removeAllocation",
         "Only the selected and flexible categories change. The target is clamped to their combined percentage. "
         "Subtracting the same delta from the flexible category preserves the exact total."),
        ("A10. Conservative auto-save validation", "data/ReceiptAutomation.kt", "fun receiptAutoSaveIssues", "fun possibleReceiptDuplicate",
         "The function collects transparent reasons for review. It does not use an unexplained percentage confidence score. "
         "Ambiguity, weak images, foreign currency and possible fuel duplication are explicit gates."),
        ("A11. Encrypted settings persistence", "data/SettingsStore.kt", "class SettingsStore", "private val settingsJson",
         "DataStoreFactory uses a custom serializer and migration. The public save method validates before updateData. "
         "The repository additionally serializes transformations so separate preferences do not overwrite a stale snapshot."),
        ("A12. Real pager and child gesture boundary", "ui/MainPager.kt", "internal data class MainDestination", "/** Static charts",
         "One list maps page keys to labels. HorizontalPager owns page dragging and snapping. "
         "Only unused horizontal child movement is consumed at the boundary. Inactive pages are visually available "
         "during a drag but do not expose duplicate focus/semantics."),
        ("A13. A preserving schema addition", "data/Database.kt", "        val MIGRATION_1_2", "        val MIGRATION_2_3",
         "ALTER TABLE adds columns with compatible defaults or nullability. Existing rows survive. "
         "A new installation and an upgraded installation must converge on the expected exported schema."),
        ("A14. Spreadsheet-safe CSV cells", "data/Backup.kt", "fun csvCell", None,
         "Formula-leading content is neutralized, quotes are doubled and every field is quoted. "
         "The UTF-8 BOM aids common spreadsheet import behavior. A CSV writer needs more than joining strings with commas."),
    ]
    for title, rel, start, end, explanation in snippets:
        path = "app/src/test/java/my/fuelledger/domain/CentsKeypadTest.kt" if rel.startswith("../") else p + rel
        parts.append(source_snippet(title, path, start, end, explanation))

    parts.append("# Appendix B. Source-file atlas\n\n"
                 "All production Kotlin files in this snapshot are listed below. Links point to the immutable "
                 "checkpoint, so they will not silently change when main advances. The offline source browser "
                 "and source ZIP contain the full files, including resources, tests and build configuration.\n\n")
    groups: dict[str, list[str]] = {}
    actual = [s.removeprefix(p) for s in paths if s.startswith(p) and s.endswith(".kt")]
    assert set(actual) == set(DESCRIPTIONS), (set(actual) - set(DESCRIPTIONS), set(DESCRIPTIONS) - set(actual))
    for rel in actual:
        groups.setdefault(rel.split("/")[0] if "/" in rel else "entry-points", []).append(rel)
    for group, items in groups.items():
        parts.append(f"## {group}\n\n")
        for rel in items:
            parts.append(f"- [{rel}]({pinned_link(p + rel)}): {DESCRIPTIONS[rel]}\n")
        parts.append("\n")
    parts.append("## Tests and evidence map\n\n"
                 "Start with the test closest to the rule you are studying. Test names describe coverage areas; "
                 "read their assertions before treating them as proof of a broader claim.\n\n")
    for prefix, label in [("app/src/test/", "Host and fixture tests"),
                          ("app/src/androidTest/", "Native Android tests"),
                          ("buildSrc/src/test/", "Build policy tests"),
                          ("scripts/tests/", "Release workflow tests")]:
        parts.append(f"### {label}\n\n")
        for path in paths:
            if path.startswith(prefix) and path.endswith((".kt", ".java", ".py")):
                short = path.removeprefix(prefix).removeprefix("java/my/fuelledger/")
                parts.append(f"- [{short}]({pinned_link(path)})\n")
        parts.append("\n")
    parts.append("## Build, resources and schemas\n\n")
    for path in ["app/build.gradle.kts", "build.gradle.kts", "buildSrc/build.gradle.kts",
                 "settings.gradle.kts", "gradle.properties", "gradle/wrapper/gradle-wrapper.properties",
                 "app/src/main/AndroidManifest.xml", "app/proguard-rules.pro", "version.properties",
                 "scripts/release.py", "scripts/setup-cloud.sh", "RELEASING.md", "SECURITY.md"]:
        if path in paths:
            parts.append(f"- [{path}]({pinned_link(path)})\n")
    parts.append("\n`app/src/main/res` contains packaged icons, themes, widget layouts and backup/provider rules. "
                 "`app/schemas` contains all five exported Room schemas. `buildSrc/src/main` contains the "
                 "dependency policy. The source ZIP preserves their directory structure.\n\n")

    parts.append("# Appendix C. Matched screens for visual study\n\n"
                 "These existing host-rendered test screenshots use the same 393×851dp, mdpi, 100% font, "
                 "Dark configuration. They are fixture captures, not personal financial records. "
                 "They demonstrate alignment and component density; they do not independently prove "
                 "calculation correctness or every physical-device behavior. See the dated CLARITY reports "
                 "for provenance and the additional light/large-font/native captures.\n\n")
    for label, key in [("Home", "Home"), ("Expenses", "Expenses"), ("Fuel", "Fuel"),
                       ("Journey", "Journey"), ("Analysis", "Analysis")]:
        source = f"docs/clarity-2026-10-18/screens/matched/393-Dark-{key}.png"
        assert source in paths, source
        dest = DOCS / "assets" / f"{label.lower()}.png"
        dest.write_bytes(git_bytes(source))
        (OUT / "assets").mkdir(exist_ok=True)
        (OUT / "assets" / dest.name).write_bytes(dest.read_bytes())
        parts.append(f"## {label}\n\n![{label}: matched synthetic test capture, Dark, 393×851dp.](assets/{dest.name})\n\n")
    parts.append("# Appendix D. Using the learning package\n\n"
                 "Read the PDF for fixed pagination and its clickable contents. Open the HTML file in a browser "
                 "for an offline reading edition with chapter navigation. The DOCX is editable in compatible "
                 "word processors; its automatic contents field may need updating after edits. Markdown is "
                 "the editable text source. The complete ZIP groups these editions and the source companion.\n\n"
                 "Open `Source_Browser.html` after extracting the package to search file names and read code "
                 "without Android Studio. Each source block retains line anchors. Extract `Source_Companion.zip` "
                 "to inspect the actual directory structure. A manifest records every included source file's "
                 "SHA-256 and the source checkpoint. No APK, signing key, personal database or machine cache "
                 "is included. First builds can require SDK/dependency downloads and a local development "
                 "signing configuration; do not substitute a generated key for the original distribution key.\n\n"
                 "The core book and source excerpts describe 1.5.4. App code and version were not changed to "
                 "produce this material. Existing app test reports are clearly historical. Documentation "
                 "validation, page count and artifact checksums are recorded with this learning edition.\n")
    return "".join(parts)


CSS = """
:root{color-scheme:light;--navy:#14283d;--green:#087757}
body{font:17px/1.65 system-ui,sans-serif;color:#18263b;max-width:880px;margin:auto;padding:30px 28px 90px}
h1{color:var(--navy);font-size:2rem;line-height:1.2;margin-top:3rem;border-top:4px solid #2abb91;padding-top:24px}
h2{font-size:1.35rem;margin-top:2.2rem}h3{font-size:1.12rem}
a{color:var(--green);text-underline-offset:3px}p,li{overflow-wrap:anywhere}
pre{background:#f1f4f6;border-left:3px solid #2abb91;padding:18px;overflow:auto;font-size:13px;line-height:1.6}
code{font-family:ui-monospace,monospace}p code,li code{background:#f1f4f6;padding:2px 4px}
table{border-collapse:collapse;width:100%;font-size:14px;margin:24px 0}
td,th{border-bottom:1px solid #ccd6dc;padding:10px;text-align:left;vertical-align:top;overflow-wrap:anywhere}
th{background:#14283d;color:white}img{max-width:100%;max-height:700px;display:block;margin:24px auto}
figcaption{font-size:13px;text-align:center;color:#536578}
#TOC{background:#f1f4f6;border-radius:16px;padding:20px;font-size:15px}
#TOC ul{padding-left:24px}header{margin-bottom:36px}.subtitle{color:#536578}
.source-line{display:block;white-space:pre}.source-line:target{background:#dcf4e9}
.ln{display:inline-block;width:4em;color:#64748b;text-align:right;margin-right:1em;text-decoration:none}
details{margin:12px 0;min-width:0}summary{cursor:pointer;font-weight:600;overflow-wrap:anywhere}
input{padding:12px;font-size:16px;width:90%;border:1px solid #94a3b8;border-radius:8px}
@media print{body{font-size:10pt;max-width:none}h1{break-before:page}pre{white-space:pre-wrap}
a{color:inherit}#TOC a{text-decoration:none}table{font-size:9pt}}
"""


def archive_sources(paths: list[str]) -> dict:
    roots = ("app/src/", "app/schemas/", "buildSrc/src/", "gradle/wrapper/", "scripts/tests/")
    explicit = {"gradlew", "gradlew.bat", ".gitignore", "settings.gradle.kts", "build.gradle.kts",
                "gradle.properties", "version.properties", "app/build.gradle.kts", "app/proguard-rules.pro",
                "buildSrc/build.gradle.kts", "scripts/release.py", "scripts/setup-cloud.sh"}
    chosen = [p for p in paths if p.startswith(roots) or p in explicit or
              ("/" not in p and p.endswith(".md")) or
              (p.startswith(("docs/journey-layout/", "docs/clarity-2026-10-18/",
                             "docs/audit-2026-10-18/")) and p.endswith((".md", ".txt", ".json")))]
    forbidden = re.compile(r"(?i)(?:^|/)(?:local\.properties|.*\.(?:jks|keystore|p12|pem|db|apk)|\.env)$")
    assert not [p for p in chosen if forbidden.search(p)]
    blocks = []
    manifest = {"source_commit": SNAPSHOT, "app_version": "1.5.4", "version_code": 11, "files": []}
    with zipfile.ZipFile(OUT / "Source_Companion.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for path in chosen:
            data = git_bytes(path)
            archive.writestr("source/" + path, data)
            manifest["files"].append({"path": path, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()})
            if path.endswith((".kt", ".kts", ".java", ".py", ".sh", ".xml", ".json", ".md", ".properties", ".pro")):
                try:
                    content = data.decode("utf-8")
                except UnicodeDecodeError:
                    continue
                anchor = hashlib.sha256(path.encode()).hexdigest()[:16]
                lines = "".join(f'<span class="source-line" id="{anchor}-L{i}"><a class="ln" href="#{anchor}-L{i}">{i}</a>{html.escape(line)}</span>'
                                for i, line in enumerate(content.splitlines(), 1))
                blocks.append(f'<details data-path="{html.escape(path.lower())}" id="{anchor}"><summary>{html.escape(path)}</summary>'
                              f'<p><a href="{pinned_link(path)}">Pinned GitHub source</a></p><pre><code>{lines}</code></pre></details>')
        archive.writestr("SOURCE_MANIFEST.json", json.dumps(manifest, indent=2) + "\n")
        archive.writestr("READ_FIRST.txt",
                         f"Source snapshot {SNAPSHOT}\nApp 1.5.4 / code 11\n"
                         "Files are exact Git blobs. No signing key, APK, financial data or build cache is included.\n"
                         "Read source/README.md, AGENTS.md, RELEASING.md and SECURITY.md.\n"
                         "Historical documentation links may refer to artifacts intentionally excluded from this source archive.\n"
                         "This is not a complete Git clone; use the repository for history and artifact downloads.\n"
                         "Local builds need the documented SDK/JDK and initial dependency downloads.\n"
                         "Preserve source notices/licenses; this package does not assign a new license to app dependencies.\n")
    browser = ("<!doctype html><html lang='en'><meta charset='utf-8'><meta name='viewport' content='width=device-width'>"
               f"<title>Fuel &amp; Expense Tracker — source companion</title><style>{CSS}</style>"
               f"<h1>Source companion</h1><p>Exact app source: 1.5.4 / code 11 · {SNAPSHOT}</p>"
               "<p>Search file paths below. Expand a file to read it; use your browser's Find for text inside expanded files. "
               "Everything needed to read this browser is embedded locally.</p>"
               "<label for='q'>Filter files</label><br><input id='q' type='search' placeholder='For example: MonthlyPlanner or Test'>"
               "<p id='count' aria-live='polite'></p>" + "".join(blocks) +
               "<script>const q=document.querySelector('#q');const fs=[...document.querySelectorAll('details')];"
               "function filter(){let n=0;for(const f of fs){f.hidden=!f.dataset.path.includes(q.value.toLowerCase());if(!f.hidden)n++}"
               "document.querySelector('#count').textContent=n+' readable files';}"
               "q.addEventListener('input',filter);filter();"
               "function reveal(){const target=document.getElementById(decodeURIComponent(location.hash.slice(1)));"
               "if(target){const d=target.closest('details');if(d){d.hidden=false;d.open=true;}target.scrollIntoView();}}"
               "addEventListener('hashchange',reveal);reveal();</script></html>")
    (OUT / "Source_Browser.html").write_text(browser)
    (DOCS / "SOURCE_MANIFEST.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


def plain(inlines: list) -> str:
    values = []
    for node in inlines:
        t, c = node["t"], node.get("c")
        if t == "Str":
            values.append(c)
        elif t in ("Space", "SoftBreak", "LineBreak"):
            values.append(" ")
        elif t in ("Code", "Math"):
            values.append(c[1])
        elif t in ("Link", "Image"):
            values.append(plain(c[1]))
        elif t in ("Strong", "Emph", "Strikeout", "SmallCaps", "Superscript", "Subscript"):
            values.append(plain(c))
        elif t == "Quoted":
            values.append('"' + plain(c[1]) + '"')
    return "".join(values)


def inline(inlines: list) -> str:
    values = []
    for node in inlines:
        t, c = node["t"], node.get("c")
        if t == "Str":
            values.append(escape(c))
        elif t in ("Space", "SoftBreak"):
            values.append(" ")
        elif t == "LineBreak":
            values.append("<br/>")
        elif t == "Code":
            values.append(f'<font face="Mono" size="8.6">{escape(c[1])}</font>')
        elif t == "Strong":
            values.append("<b>" + inline(c) + "</b>")
        elif t == "Emph":
            values.append("<i>" + inline(c) + "</i>")
        elif t == "Link":
            values.append(f'<a href="{escape(c[2][0], {chr(34): "&quot;"})}" color="#087757">{inline(c[1])}</a>')
        elif t == "Quoted":
            values.append("“" + inline(c[1]) + "”")
        elif t == "Image":
            values.append(inline(c[1]))
        elif t in ("Superscript", "Subscript", "Strikeout", "SmallCaps"):
            values.append(inline(c))
        else:
            raise ValueError(f"Unsupported inline: {t}")
    return "".join(values)


def make_pdf(ast: dict, destination: Path) -> None:
    font_root = Path("/usr/share/fonts/truetype")
    for name, path in [
        ("Body", "open-sans/OpenSans-Regular.ttf"), ("Body-Bold", "open-sans/OpenSans-Bold.ttf"),
        ("Body-Italic", "open-sans/OpenSans-Italic.ttf"), ("Body-BoldItalic", "open-sans/OpenSans-BoldItalic.ttf"),
        ("Mono", "liberation/LiberationMono-Regular.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(font_root / path)))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-Bold", italic="Body-Italic", boldItalic="Body-BoldItalic")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("BodyBook", fontName="Body", fontSize=10.1, leading=15.4, spaceAfter=9,
                             textColor=colors.HexColor("#18263b"), splitLongWords=True))
    styles.add(ParagraphStyle("SmallBook", parent=styles["BodyBook"], fontSize=8.1, leading=11.5, spaceAfter=5))
    styles.add(ParagraphStyle("CellBook", parent=styles["BodyBook"], fontSize=8.2, leading=12, spaceAfter=0))
    styles.add(ParagraphStyle("CellHead", parent=styles["CellBook"], fontName="Body-Bold", textColor=colors.white))
    styles.add(ParagraphStyle("CodeBook", fontName="Mono", fontSize=7.5, leading=10.7,
                             leftIndent=8, rightIndent=5, spaceBefore=6, spaceAfter=12,
                             backColor=colors.HexColor("#f0f4f6"), textColor=colors.HexColor("#14283d")))
    for level, size, leading in [(1, 23, 30), (2, 14.5, 20), (3, 11.5, 16)]:
        styles.add(ParagraphStyle(f"H{level}Book", fontName="Body-Bold", fontSize=size, leading=leading,
                                 textColor=colors.HexColor("#14283d"), spaceBefore=16 if level > 1 else 0,
                                 spaceAfter=12, keepWithNext=True))
    styles.add(ParagraphStyle("CoverTitle", fontName="Body-Bold", fontSize=33, leading=42,
                             textColor=colors.white, spaceAfter=24))
    styles.add(ParagraphStyle("CoverText", fontName="Body", fontSize=12, leading=19,
                             textColor=colors.HexColor("#dce8ef"), spaceAfter=12))
    styles.add(ParagraphStyle("TOCBook", fontName="Body", fontSize=9.3, leading=14.5,
                             leftIndent=0, firstLineIndent=0, spaceBefore=4))
    width, height = A4
    body_width = width - 96

    class Book(BaseDocTemplate):
        def afterFlowable(self, flowable):
            if hasattr(flowable, "book_heading"):
                level, title, anchor = flowable.book_heading
                self.canv.bookmarkPage(anchor)
                self.canv.addOutlineEntry(title, anchor, level=level - 1, closed=level > 1)
                if level == 1:
                    self.notify("TOCEntry", (0, title, self.page, anchor))

    def background(canvas, doc):
        if doc.page == 1:
            canvas.setFillColor(colors.HexColor("#0B1220"))
            canvas.rect(0, 0, width, height, fill=1, stroke=0)
            canvas.setFillColor(colors.HexColor("#2abb91"))
            canvas.rect(48, height - 95, 95, 6, fill=1, stroke=0)
            canvas.setFillColor(colors.HexColor("#94a3b8"))
            canvas.setFont("Body", 9)
            canvas.drawString(48, 50, "A code-based course in building and maintaining an offline Android app")
        else:
            canvas.setStrokeColor(colors.HexColor("#ccd6dc"))
            canvas.line(48, height - 32, width - 48, height - 32)
            canvas.setFont("Body", 8)
            canvas.setFillColor(colors.HexColor("#536578"))
            canvas.drawString(48, height - 24, "FUEL & EXPENSE TRACKER  /  LEARNING GUIDE")
            canvas.drawString(48, 25, "Source 1.5.4 · Learning edition 21 October 2026")
            canvas.drawRightString(width - 48, 25, str(doc.page))

    doc = Book(str(destination), pagesize=A4, leftMargin=48, rightMargin=48, topMargin=52, bottomMargin=45,
               title=TITLE, author="Fuel & Expense Tracker learning material")
    doc.addPageTemplates(PageTemplate(id="book", frames=[Frame(48, 45, body_width, height - 97,
                                                               leftPadding=0, rightPadding=0,
                                                               topPadding=0, bottomPadding=0)], onPage=background))
    story = [
        Spacer(1, 85),
        Paragraph("Fuel &amp; Expense<br/>Tracker", styles["CoverTitle"]),
        Paragraph("Learn Android development<br/>through your own application", styles["CoverText"]),
        Spacer(1, 30),
        Paragraph("Kotlin · Compose · exact money · Room<br/>fuel and journey logic · offline receipts<br/>privacy · testing · GitHub · safe releases", styles["CoverText"]),
        Spacer(1, 30),
        Paragraph("29 chapters · worked examples · guided labs<br/>annotated source excerpts · complete source companion", styles["CoverText"]),
        Spacer(1, 20),
        Paragraph("Learning edition: 21 October 2026<br/>App source: 1.5.4 / version code 11", styles["CoverText"]),
        PageBreak(),
        Paragraph("Contents", styles["H1Book"]),
        Paragraph("Follow the numbered chapters for the course. Use the appendices while reading the real code. "
                  "Contents entries and source references are clickable.", styles["BodyBook"]),
    ]
    toc = TableOfContents()
    toc.levelStyles = [styles["TOCBook"]]
    story.extend([toc, PageBreak()])
    seen_heading = False
    bookmarks = 0

    def paragraph_blocks(blocks):
        return "<br/>".join(inline(b["c"]) for b in blocks if b["t"] in ("Para", "Plain"))

    def convert(blocks, nesting=0):
        nonlocal seen_heading, bookmarks
        result = []
        for block in blocks:
            t, c = block["t"], block.get("c")
            if t == "Header":
                level = min(c[0], 3)
                if level == 1 and seen_heading:
                    result.append(PageBreak())
                seen_heading = True
                para = Paragraph(inline(c[2]), styles[f"H{level}Book"])
                bookmarks += 1
                para.book_heading = (level, plain(c[2]), f"heading-{bookmarks}")
                result.append(para)
            elif t in ("Para", "Plain"):
                if len(c) == 1 and c[0]["t"] == "Image":
                    image_node = c[0]["c"]
                    path = DOCS / image_node[2][0]
                    pic = Image(str(path))
                    ratio = min(body_width / pic.imageWidth, 510 / pic.imageHeight)
                    pic.drawWidth, pic.drawHeight = pic.imageWidth * ratio, pic.imageHeight * ratio
                    pic.hAlign = "CENTER"
                    preceding = []
                    if result and hasattr(result[-1], "book_heading"):
                        preceding.append(result.pop())
                    result.append(KeepTogether(preceding + [pic, Spacer(1, 8),
                                                Paragraph(inline(image_node[1]), styles["SmallBook"])]))
                else:
                    result.append(Paragraph(inline(c), styles["BodyBook"]))
            elif t == "CodeBlock":
                wrapped = []
                for line in c[1].expandtabs(4).splitlines():
                    # Visual wrapping only; original exact bytes stay in Markdown and source companion.
                    wrapped.extend(textwrap.wrap(line, width=105, break_long_words=True,
                                                 break_on_hyphens=False, replace_whitespace=False,
                                                 drop_whitespace=False, subsequent_indent="    ") or [""])
                result.append(Preformatted("\n".join(wrapped), styles["CodeBook"]))
            elif t in ("BulletList", "OrderedList"):
                items = c if t == "BulletList" else c[1]
                start = 1 if t == "BulletList" else c[0][0]
                for index, item in enumerate(items):
                    marker = "•" if t == "BulletList" else f"{start + index}."
                    result.append(Paragraph(f"{marker}  {paragraph_blocks(item)}",
                                            ParagraphStyle("item", parent=styles["BodyBook"], leftIndent=12,
                                                           firstLineIndent=-10, spaceAfter=6)))
            elif t == "Table":
                # Pandoc 3 table structure: attr, caption, colspecs, head, bodies, foot.
                head_rows = c[3][1]
                body_rows = [r for body in c[4] for r in body[2] + body[3]]
                rows = head_rows + body_rows + c[5][1]
                grid = []
                for row_index, row in enumerate(rows):
                    style = styles["CellHead"] if row_index < len(head_rows) else styles["CellBook"]
                    grid.append([Paragraph(paragraph_blocks(cell[4]), style) for cell in row[1]])
                cols = len(c[2])
                fractions = {2: [0.31, 0.69], 3: [0.25, 0.34, 0.41], 4: [0.20, 0.22, 0.27, 0.31]}.get(cols, [1/cols] * cols)
                table = Table(grid, colWidths=[body_width*x for x in fractions], repeatRows=len(head_rows), hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, len(head_rows)-1), colors.HexColor("#14283d")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccd6dc")),
                    ("ROWBACKGROUNDS", (0, len(head_rows)), (-1, -1), [colors.white, colors.HexColor("#f5f7f8")]),
                ]))
                result.extend([table, Spacer(1, 12)])
            elif t == "BlockQuote":
                result.extend(convert(c, nesting+1))
            elif t == "HorizontalRule":
                result.append(Spacer(1, 14))
            else:
                raise ValueError(f"Unsupported block: {t}")
        return result

    story.extend(convert(ast["blocks"]))
    doc.multiBuild(story)


def main():
    OUT.mkdir(exist_ok=True)
    (DOCS / "assets").mkdir(exist_ok=True)
    paths = subprocess.check_output(["git", "ls-tree", "-r", "--name-only", SNAPSHOT], cwd=ROOT, text=True).splitlines()
    chapters = sorted((DOCS / "chapters").glob("*.md"))
    core = "\n\n".join(path.read_text() for path in chapters)
    assert [int(x) for x in re.findall(r"^# (\d+)\.", core, flags=re.M)] == list(range(1, 30))
    appendix = appendices(paths)
    (DOCS / "APPENDICES.md").write_text(appendix)
    combined = core + "\n\n" + appendix
    (DOCS / "BOOK.md").write_text(combined)
    (OUT / f"{BASE}.md").write_text(combined)
    (DOCS / "assets/reading.css").write_text(CSS)
    subprocess.run(["pandoc", "BOOK.md", "-f", "gfm", "-t", "html5", "--standalone", "--toc", "--toc-depth=2",
                    "--embed-resources", "--css=assets/reading.css", "--metadata", f"title={TITLE}",
                    "--metadata", "subtitle=29 chapters · source 1.5.4 · 21 October 2026",
                    "-o", str(OUT / f"{BASE}.html")], cwd=DOCS, check=True)
    html_path = OUT / f"{BASE}.html"
    html_path.write_text("\n".join(line.rstrip() for line in html_path.read_text().splitlines()) + "\n")
    subprocess.run(["pandoc", "BOOK.md", "-f", "gfm", "--toc", "--toc-depth=1",
                    "--metadata", f"title={TITLE}", "-o", str(OUT / f"{BASE}.docx")], cwd=DOCS, check=True)
    docx = Document(OUT / f"{BASE}.docx")
    for section in docx.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.left_margin = section.right_margin = Inches(.8)
    docx.styles["Normal"].font.name = "Arial"
    docx.styles["Normal"].font.size = Pt(10.5)
    for p in docx.paragraphs:
        if p.style.name == "Heading 1":
            p.paragraph_format.page_break_before = True
    for picture in docx.inline_shapes:
        if picture.height > Inches(7):
            ratio = Inches(7) / picture.height
            picture.width, picture.height = int(picture.width * ratio), int(picture.height * ratio)
    for table in docx.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    docx.save(OUT / f"{BASE}.docx")
    ast = json.loads(subprocess.check_output(["pandoc", "BOOK.md", "-f", "gfm", "-t", "json"], cwd=DOCS))
    make_pdf(ast, OUT / f"{BASE}.pdf")
    manifest = archive_sources(paths)
    (OUT / "READ_FIRST.txt").write_text(
        "Fuel & Expense Tracker learning edition — 21 October 2026\n"
        "Start with Fuel_Expense_Tracker_Learning_Guide.pdf.\n"
        "HTML reads offline; DOCX and Markdown are editable.\n"
        "Source_Browser.html is an offline file browser; Source_Companion.zip contains exact source.\n"
        "Source checkpoint: " + SNAPSHOT + "\n"
        "App source 1.5.4 / code 11. No APK release or application changes were made for this book.\n"
        "The source archive excludes secrets, user records, APKs and build caches.\n"
        "External reference links require connectivity; the reading editions and source browser do not.\n"
        "To rebuild these documents in the repository, run docs/learning/tools/build_book.py.\n")
    stats = {
        "source_commit": SNAPSHOT, "app_version": "1.5.4", "version_code": 11,
        "learning_edition": "2026-10-21", "chapters": 29,
        "core_word_count": len(core.split()), "combined_word_count": len(combined.split()),
        "pdf_pages": len(PdfReader(OUT / f"{BASE}.pdf").pages),
        "source_files": len(manifest["files"]), "production_kotlin_files": len(DESCRIPTIONS),
    }
    (DOCS / "BUILD_STATS.json").write_text(json.dumps(stats, indent=2) + "\n")
    members = [f"{BASE}.{ext}" for ext in ("pdf", "docx", "html", "md")] + [
        "Source_Browser.html", "Source_Companion.zip", "READ_FIRST.txt"]
    with zipfile.ZipFile(OUT / "Fuel_Expense_Tracker_Complete_Learning_Pack.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for name in members:
            archive.write(OUT / name, name)
        for image in sorted((DOCS / "assets").glob("*.png")):
            archive.write(image, f"assets/{image.name}")
        archive.writestr("BUILD_STATS.json", json.dumps(stats, indent=2) + "\n")
    artifacts = [*members, "Fuel_Expense_Tracker_Complete_Learning_Pack.zip"]
    (OUT / "SHA256SUMS").write_text("".join(f"{hashlib.sha256((OUT/name).read_bytes()).hexdigest()}  {name}\n"
                                         for name in artifacts))
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()
